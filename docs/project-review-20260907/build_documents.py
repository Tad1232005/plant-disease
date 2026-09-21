"""Build the Vietnamese project document pack with the bundled Python runtime."""
from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'word'
OUT.mkdir(exist_ok=True)


def text_runs(p, text):
    for i, part in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)', text)):
        if not part:
            continue
        r = p.add_run(part[2:-2] if part.startswith('**') else part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            r.bold = True
        if part.startswith('`'):
            r.font.name = 'Consolas'
            r.font.size = Pt(9)


def table(doc, rows):
    headers = rows[0]
    n = len(headers)
    t = doc.add_table(rows=1, cols=n)
    t.style = 'Table Grid'
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Intentional widths for endpoint records and dictionaries.
    if headers[0] == 'Method':
        widths = [0.65, 2.40, 1.02, 2.93] if n == 4 else [0.65, 2.60, 1.05, 1.60, 1.10]
    elif headers[0] == 'Cột hoặc nhóm cột':
        widths = [1.90, 2.25, 2.85]
    elif headers[0] == 'Ngày':
        widths = [0.50, 3.00, 3.50]
    elif headers[0] == 'Mã':
        widths = [0.55, 2.75, 3.70]
    elif headers[0] == 'Rủi ro':
        widths = [1.95, 0.50, 1.00, 3.55]
    elif n == 2:
        widths = [2.10, 4.90]
    elif n == 3:
        widths = [1.45, 2.55, 3.00]
    elif n == 4:
        widths = [1.15, 1.50, 1.85, 2.50]
    elif n == 5:
        widths = [1.15, 1.00, 1.20, 1.30, 2.35]
    else:
        widths = [1.7] + [(7.0-1.7)/(n-1)]*(n-1)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for idx, row in enumerate(rows):
        cells = t.rows[0].cells if idx == 0 else t.add_row().cells
        for j, (c, value) in enumerate(zip(cells, row)):
            c.width = Inches(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E3EBF1' if idx == 0 else ('F5F7F9' if idx % 2 == 0 else 'FFFFFF'))
            pr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m = OxmlElement('w:'+side)
                m.set(qn('w:w'),'50' if headers[0]=='Cột hoặc nhóm cột' and side in ['top','bottom'] else '85')
                m.set(qn('w:type'),'dxa')
                margins.append(m)
            pr.append(margins)
            p = c.paragraphs[0]
            if (j == 0 and headers[0] in ['Method','Ngày','Mã']) or (headers[0]=='Chức năng' and j>0):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.06
            text_runs(p, value.replace('<br>', '\n'))
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0,0,0)
                if idx == 0: r.bold = True
        trpr = t.rows[idx]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement('w:'+edge)
        b.set(qn('w:val'),'single')
        b.set(qn('w:sz'),'4')
        b.set(qn('w:color'),'D9D9D9')
        borders.append(b)
    t._tbl.tblPr.append(borders)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(6)
    spacer.add_run().font.size = Pt(1)


def build(src):
    doc = Document()
    # Strip inherited decorative paragraph rules from the bundled default.
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = Inches(0.66)
    sec.bottom_margin = Inches(0.65)
    sec.footer_distance = Inches(0.25)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Calibri')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [('Title',24),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
        s=doc.styles[name]
        s.font.size=Pt(size)
        s.font.bold=True
        s.paragraph_format.space_before=Pt(10 if name != 'Title' else 0)
        s.paragraph_format.space_after=Pt(7)
        s.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Plant Disease  |  ')
    r.font.size=Pt(8)
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    doc.core_properties.author='Plant Disease Project'
    doc.core_properties.subject='Đặc tả và kế hoạch hiệu chỉnh dự án 8 tuần'
    lines=src.read_text(encoding='utf-8').splitlines()
    idx=0
    page_break_pending=False
    while idx<len(lines):
        line=lines[idx].strip()
        if not line:
            idx+=1;continue
        if line=='[PAGE]':
            page_break_pending=True
        elif line.startswith('|'):
            rows=[]
            while idx<len(lines) and lines[idx].strip().startswith('|'):
                row=[x.strip() for x in lines[idx].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x or 'x') for x in row): rows.append(row)
                idx+=1
            if any(len(row)!=len(rows[0]) for row in rows):raise ValueError(rows)
            table(doc, rows)
            continue
        elif line.startswith('```'):
            idx+=1
            while idx<len(lines) and not lines[idx].startswith('```'):
                p=doc.add_paragraph()
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.02
                r=p.add_run(lines[idx]);r.font.name='Consolas';r.font.size=Pt(9)
                idx+=1
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
        elif line.startswith('# '):
            p=doc.add_paragraph(line[2:],style='Title')
            doc.core_properties.title=line[2:]
        elif line.startswith('### '):
            doc.add_paragraph(line[4:].replace('_',' '),style='Heading 2')
        elif line.startswith('## '):
            p=doc.add_paragraph(line[3:],style='Heading 1')
            if page_break_pending:
                p.paragraph_format.page_break_before=True
                page_break_pending=False
        elif line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet');text_runs(p,line[2:])
        else:
            p=doc.add_paragraph();text_runs(p,line)
        idx+=1
    path=OUT/(src.stem+'.docx')
    doc.save(path)
    return str(path)


if __name__=='__main__':
    print(json.dumps([build(p) for p in sorted(ROOT.glob('0*.md'))],ensure_ascii=False,indent=2))
